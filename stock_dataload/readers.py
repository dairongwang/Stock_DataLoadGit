"""
Source file readers: CSV / delimited text, Excel (.xlsx), JSON and XML.

Every reader returns a :class:`SourceTable` - an optional header plus an
iterator of rows where every cell is already a string.  The converter does
not care where the table came from; mapping/validation/writing is shared.
"""

from __future__ import annotations

import csv
import json
import logging
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Iterator

from .file_formats import FileFormat

log = logging.getLogger(__name__)

CANDIDATE_DELIMITERS = (",", "\t", ";", "|")

EXTENSION_TYPES = {
    ".csv": "csv",
    ".txt": "text", ".tsv": "text", ".dat": "text", ".prn": "text", ".psv": "text",
    ".xlsx": "excel", ".xlsm": "excel",
    ".json": "json", ".jsonl": "json", ".ndjson": "json",
    ".xml": "xml",
}
SUPPORTED_EXTENSIONS = tuple(EXTENSION_TYPES) + (".xls",)


class ReadError(Exception):
    """The source file could not be read/understood."""


@dataclass
class SourceTable:
    file_type: str
    header: list[str] | None
    rows: Iterator[list[str]]
    delimiter: str | None = None       # only for delimited files
    detail: str = ""                   # human readable note (sheet name, record path, ...)


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def to_text(value: Any) -> str:
    """Render a cell value (from JSON/Excel) the way it would appear in a text file."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == 0:
            return value.strftime("%Y%m%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y%m%d")
    if isinstance(value, (list, dict)):
        return json.dumps(value, separators=(",", ":"))
    return str(value).strip()


def detect_delimiter(sample_lines: Iterable[str]) -> str:
    """Pick the candidate delimiter that appears consistently on the lines."""
    best, best_score = ",", -1
    lines = [ln for ln in sample_lines if ln.strip()]
    if not lines:
        return best
    for d in CANDIDATE_DELIMITERS:
        counts = [ln.count(d) for ln in lines]
        if min(counts) == 0:
            continue
        consistent = len(set(counts)) == 1
        score = (2 if consistent else 0) * 1000 + min(counts)
        if score > best_score:
            best, best_score = d, score
    return best


def detect_file_type(path: Path, declared: str = "auto") -> str:
    """Resolve 'auto' to csv/text/excel/json/xml using extension, then content."""
    if declared and declared != "auto":
        return declared
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".xls":
        raise ReadError("legacy .xls workbooks are not supported - save the file as .xlsx")
    if ext in EXTENSION_TYPES:
        kind = EXTENSION_TYPES[ext]
        if kind in ("csv", "text"):
            # a .txt could still be json/xml - peek
            sniffed = _sniff(path)
            return sniffed if sniffed in ("json", "xml") else kind
        return kind
    sniffed = _sniff(path)
    return sniffed or "text"


def _sniff(path: Path) -> str | None:
    try:
        with open(path, "rb") as fh:
            head = fh.read(2048)
    except OSError:
        return None
    if head.startswith(b"PK\x03\x04"):
        return "excel"
    text = head.lstrip(b"\xef\xbb\xbf").lstrip()
    first_line = text.split(b"\n", 1)[0]
    if text.startswith((b"<?xml", b"<!--", b"<!DOCTYPE")):
        return "xml"
    if text.startswith(b"<") and b">,<" not in first_line and b">\t<" not in first_line \
            and b">;<" not in first_line and b">|<" not in first_line:
        # looks like a root element rather than a "<ticker>,<date>,..." header
        return "xml"
    if text.startswith(b"{") or text.startswith(b"["):
        return "json"
    return None


def _flatten(obj: Any, prefix: str = "", out: dict[str, Any] | None = None) -> dict[str, Any]:
    """Flatten nested dicts: {"price": {"open": 1}} -> {"price.open": 1}."""
    out = {} if out is None else out
    if isinstance(obj, dict):
        for k, v in obj.items():
            key = f"{prefix}.{k}" if prefix else str(k)
            if isinstance(v, dict):
                _flatten(v, key, out)
            else:
                out[key] = v
    return out


def _dict_records_to_table(records: list[dict[str, Any]]) -> tuple[list[str], list[list[str]]]:
    flat = [_flatten(r) for r in records]
    header: list[str] = []
    seen: set[str] = set()
    for rec in flat:
        for k in rec:
            if k not in seen:
                seen.add(k)
                header.append(k)
    rows = [[to_text(rec.get(k)) for k in header] for rec in flat]
    return header, rows


# ---------------------------------------------------------------------------
# delimited text / csv
# ---------------------------------------------------------------------------
def read_delimited(path: Path, fmt: FileFormat) -> SourceTable:
    fh = open(path, "r", newline="", encoding=fmt.encoding)
    try:
        pos = fh.tell()
        sample: list[str] = []
        for _ in range(fmt.skip_rows):
            fh.readline()
        data_start = fh.tell()
        for _ in range(20):
            line = fh.readline()
            if not line:
                break
            sample.append(line)
        fh.seek(data_start)
        delimiter = fmt.delimiter or detect_delimiter(sample)
        reader = csv.reader(fh, delimiter=delimiter, skipinitialspace=True)
        header = None
        if fmt.has_header:
            header = [h.strip() for h in next(reader, [])]

        def rows() -> Iterator[list[str]]:
            try:
                for row in reader:
                    yield [c.strip() for c in row]
            finally:
                fh.close()
        del pos
        return SourceTable("csv" if delimiter == "," else "text", header, rows(), delimiter=delimiter)
    except Exception:
        fh.close()
        raise


# ---------------------------------------------------------------------------
# excel (.xlsx)
# ---------------------------------------------------------------------------
def read_excel(path: Path, fmt: FileFormat) -> SourceTable:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise ReadError("reading Excel files requires the openpyxl package: pip install openpyxl") from exc
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ReadError(f"cannot open workbook {path.name}: {exc}") from exc
    if fmt.sheet:
        if fmt.sheet.isdigit():
            idx = int(fmt.sheet) - 1
            if idx < 0 or idx >= len(wb.sheetnames):
                raise ReadError(f"workbook has {len(wb.sheetnames)} sheet(s); sheet {fmt.sheet} does not exist")
            ws = wb.worksheets[idx]
        elif fmt.sheet in wb.sheetnames:
            ws = wb[fmt.sheet]
        else:
            raise ReadError(f"sheet '{fmt.sheet}' not found; available: {wb.sheetnames}")
    else:
        ws = wb.worksheets[0]

    it = ws.iter_rows(values_only=True)
    for _ in range(fmt.skip_rows):
        next(it, None)

    def non_empty(row) -> bool:
        return any(v is not None and str(v).strip() != "" for v in row)

    header = None
    if fmt.has_header:
        for row in it:
            if non_empty(row):
                header = [to_text(v) for v in row]
                break
        if header is None:
            header = []
        while header and header[-1] == "":
            header.pop()

    def rows() -> Iterator[list[str]]:
        try:
            for row in it:
                if non_empty(row):
                    yield [to_text(v) for v in row]
        finally:
            wb.close()

    return SourceTable("excel", header, rows(), detail=f"sheet '{ws.title}'")


# ---------------------------------------------------------------------------
# json
# ---------------------------------------------------------------------------
def _walk_path(obj: Any, path: str) -> Any:
    for part in [p for p in path.replace("/", ".").split(".") if p]:
        if isinstance(obj, dict):
            if part not in obj:
                raise ReadError(f"record_path '{path}': key '{part}' not found")
            obj = obj[part]
        elif isinstance(obj, list) and part.isdigit():
            obj = obj[int(part)]
        else:
            raise ReadError(f"record_path '{path}': cannot descend into '{part}'")
    return obj


def _find_records(obj: Any, depth: int = 0) -> tuple[str, list] | None:
    """Find the first list of records (dicts or lists) in a JSON structure."""
    if isinstance(obj, list):
        if obj and all(isinstance(x, (dict, list)) for x in obj):
            return "", obj
        return None
    if isinstance(obj, dict) and depth < 6:
        for k, v in obj.items():
            found = _find_records(v, depth + 1)
            if found is not None:
                sub, recs = found
                return (f"{k}.{sub}" if sub else k), recs
    return None


def read_json(path: Path, fmt: FileFormat) -> SourceTable:
    text = Path(path).read_text(encoding=fmt.encoding)
    stripped = text.strip()
    try:
        if path.suffix.lower() in (".jsonl", ".ndjson") or (stripped.startswith("{") and "\n{" in stripped):
            data: Any = [json.loads(ln) for ln in stripped.splitlines() if ln.strip()]
        else:
            data = json.loads(stripped)
    except json.JSONDecodeError as exc:
        raise ReadError(f"{path.name} is not valid JSON: {exc}") from exc

    detail = ""
    if fmt.record_path:
        records = _walk_path(data, fmt.record_path)
        detail = f"records at '{fmt.record_path}'"
    else:
        found = _find_records(data)
        if found is None:
            raise ReadError("could not find a list of records in the JSON document - set record_path")
        detail, records = found
        detail = f"records at '{detail or '$'}'"
    if not isinstance(records, list):
        raise ReadError("record_path must point to a list of records")
    # column-oriented: {"ticker": [...], "close": [...]} ?
    if isinstance(records, dict):
        raise ReadError("record_path must point to a list")

    if records and isinstance(records[0], dict):
        header, rows = _dict_records_to_table([r for r in records if isinstance(r, dict)])
        return SourceTable("json", header, iter(rows), detail=detail)
    # list of lists: optional header row
    rows = [[to_text(v) for v in r] for r in records if isinstance(r, list)]
    header = None
    if fmt.has_header and rows:
        header, rows = rows[0], rows[1:]
    return SourceTable("json", header, iter(rows), detail=detail + " (array rows)")


# ---------------------------------------------------------------------------
# xml
# ---------------------------------------------------------------------------
def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _guess_record_tag(root: ET.Element) -> str:
    """Most frequent element whose children are all leaves (a 'record'), e.g. <quote>/<row>.

    Falls back to the most frequent childless element with attributes
    (<row symbol=".." close=".."/> style).
    """
    record_counts: dict[str, int] = {}
    attr_counts: dict[str, int] = {}
    order: list[str] = []
    for el in root.iter():
        if el is root:
            continue
        t = _local(el.tag)
        if t not in order:
            order.append(t)
        if len(el) and all(len(child) == 0 for child in el):
            record_counts[t] = record_counts.get(t, 0) + 1
        elif not len(el) and el.attrib:
            attr_counts[t] = attr_counts.get(t, 0) + 1
    for counts in (record_counts, attr_counts):
        if counts:
            return max((t for t in order if t in counts), key=lambda t: counts[t])
    raise ReadError("no record elements found in XML - set record_tag")


def read_xml(path: Path, fmt: FileFormat) -> SourceTable:
    try:
        tree = ET.parse(path)
    except ET.ParseError as exc:
        raise ReadError(f"{path.name} is not well-formed XML: {exc}") from exc
    root = tree.getroot()
    tag = fmt.record_tag.strip() or _guess_record_tag(root)
    elements = [el for el in root.iter() if _local(el.tag) == tag]
    if not elements:
        raise ReadError(f"no <{tag}> elements found in {path.name}")
    records: list[dict[str, Any]] = []
    for el in elements:
        rec: dict[str, Any] = {}
        for k, v in el.attrib.items():
            rec[_local(k)] = v
        for child in el:
            name = _local(child.tag)
            # <field name="ticker">AAPL</field> style
            if name in ("field", "col", "column", "value") and ("name" in child.attrib or "id" in child.attrib):
                name = child.attrib.get("name") or child.attrib.get("id")
            if len(child):
                for sub in child:
                    rec[f"{name}.{_local(sub.tag)}"] = (sub.text or "").strip()
            else:
                rec[name] = (child.text or "").strip()
        if not rec and (el.text or "").strip():
            rec["value"] = el.text.strip()
        records.append(rec)
    header, rows = _dict_records_to_table(records)
    return SourceTable("xml", header, iter(rows), detail=f"<{tag}> elements ({len(rows)})")


# ---------------------------------------------------------------------------
READERS = {
    "csv": read_delimited,
    "text": read_delimited,
    "excel": read_excel,
    "json": read_json,
    "xml": read_xml,
}


def read_source(path: Path, fmt: FileFormat) -> SourceTable:
    """Open *path* according to *fmt* and return a SourceTable."""
    path = Path(path)
    if not path.is_file():
        raise ReadError(f"source file not found: {path}")
    kind = detect_file_type(path, fmt.file_type)
    reader = READERS.get(kind)
    if reader is None:
        raise ReadError(f"unsupported file type '{kind}'")
    table = reader(path, fmt)
    if kind in ("csv", "text") and table.file_type != kind:
        table.file_type = kind if fmt.file_type != "auto" else table.file_type
    return table
